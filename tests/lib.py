"""src/lib on its own: tests/lib.mjs cleans every fixture in Node, this asserts the cleaned bytes.
No page, no browser. Run: python3 tests/lib.py
"""
import io, json, os, re, subprocess, sys, zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'out', 'lib')
CASES = ['photo.jpg', 'photo.heic', 'c2pa.heic', 'photo.webp', 'video.mp4', 'video.mov', 'memo.m4a', 'song.mp3', 'fees.odt', 'fees.ods', 'fees.odp', 'shot.png', 'c2pa.jpg', 'c2pa.png', 'offer.docx', 'fees.xlsx', 'pitch.pptx', 'memo.pdf', 'memo-z.pdf', 'objstm.pdf', 'qpdf.pdf']
failures = []


def check(cond, msg):
    print(('  ok   ' if cond else '  FAIL ') + msg)
    if not cond:
        failures.append(msg)


LEAKS = [b'Villanueva', b'Priya', b'Delacroix', b'Lightroom', b'GIMP', b'iPhone', b'Quezon', b'Tamarind', b'DEL-2026', b'Byrne', b'confidential', b'c2pa', b'C2PA Test Signing Cert', b'14.6760', b'Mabini', b'Gate code', b'LaserJet', b'Lavf', b'Jane Secret', b'004A0061006E0065', b'Brien', b'Okafor', b'Leeds']

failures = []


def check(cond, msg):
    print(('  ok   ' if cond else '  FAIL ') + msg)
    if not cond:
        failures.append(msg)


def verify(name, data, before, fig=''):
    ext = name.rsplit('.', 1)[1]
    leaks = [l for l in LEAKS if l in data]
    if ext == 'pdf':
        # Same length, blanked in place: the words must be gone but the structure untouched.
        check(len(data) == len(before), 'pdf: same length as the original')
        check(not leaks, f'pdf: no leaked strings ({leaks})')
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(data))
        check(len(r.pages) == 1 and 'Fee summary' in r.pages[0].extract_text(), 'pdf: opens and the page text is intact')
        info = {k: str(v).strip() for k, v in (r.metadata or {}).items()}
        check(all(not v for v in info.values()), f'pdf: Info values blank ({info})')
        xmp = r.xmp_metadata
        blank = lambda v: not ''.join(v if isinstance(v, list) else [v or '']).strip()
        # Compressed XMP and object streams are rewritten to the same compressed length, so nothing moves.
        check(xmp is None or (blank(xmp.dc_creator) and blank(xmp.xmp_creator_tool)), f'{name}: XMP creator and tool blank ({xmp and xmp.dc_creator}, {xmp and xmp.xmp_creator_tool})')
        check(fig['kept'] == 0 and 'shown and kept' not in fig['note'] and 'not cleaned' not in fig['note'], f'{name}: nothing kept, the note claims no gap ({fig})')
        tmp = os.path.join(OUT, 'probe.pdf')
        open(tmp, 'wb').write(data)
        gs = subprocess.run(['gs', '-q', '-dNOPAUSE', '-dBATCH', '-sDEVICE=nullpage', tmp], capture_output=True, text=True)
        check(gs.returncode == 0 and not gs.stdout.strip() and not gs.stderr.strip(), f'{name}: Ghostscript renders it without a warning ({(gs.stdout + gs.stderr)[:200]})')
        import pikepdf
        problems = pikepdf.open(tmp).check_pdf_syntax()
        check(not problems, f'{name}: qpdf finds no syntax problems ({problems})')
        return
    if ext in ('mp4', 'mov', 'm4a'):
        check(len(data) == len(before), f'{ext}: same length, cleaned in place')
        check(not leaks, f'{ext}: no leaked strings anywhere in the bytes ({leaks})')
        tmp = os.path.join(OUT, 'probe.' + ext)
        open(tmp, 'wb').write(data)
        probe = lambda p: subprocess.run(['ffprobe', '-v', 'error', '-show_entries', 'format_tags:stream_tags', '-of', 'json', p], capture_output=True, text=True).stdout
        tags = probe(tmp)
        check(not any(t in tags for t in ('location', 'ISO6709', 'make', 'model', 'Gate code', '2026-04-11')), f'{ext}: no location, camera, title or date tags ({tags})')
        frames = lambda p: subprocess.run(['ffmpeg', '-v', 'error', '-i', p, '-f', 'framemd5', '-'], capture_output=True, text=True)
        a, b = frames(os.path.join(HERE, 'files', name)), frames(tmp)
        check(not b.stderr and a.stdout == b.stdout, f'{ext}: decodes without errors, frames and sound identical ({b.stderr[:200]})')
        return
    if ext == 'mp3':
        from mutagen.mp3 import MP3
        check(not leaks, f'mp3: no leaked strings ({leaks})')
        check(not MP3(io.BytesIO(data)).tags and b'TAG' != data[-128:-125] and not data.startswith(b'ID3'), 'mp3: no ID3v2 or ID3v1 tag left')
        tmp = os.path.join(OUT, 'probe.mp3')
        open(tmp, 'wb').write(data)
        frames = lambda p: subprocess.run(['ffmpeg', '-v', 'error', '-i', p, '-map', '0:a', '-f', 'framemd5', '-'], capture_output=True, text=True)
        a, b = frames(os.path.join(HERE, 'files', name)), frames(tmp)
        check(not b.stderr and a.stdout == b.stdout and a.stdout.count('\n') > 20, f'mp3: decodes, sound identical ({b.stderr[:200]})')
        return
    if ext == 'webp':
        from PIL import Image
        check(not leaks, f'webp: no leaked strings ({leaks})')
        im, orig = Image.open(io.BytesIO(data)), Image.open(io.BytesIO(before))
        check(im.tobytes() == orig.tobytes(), f'webp: Pillow opens it, picture identical {im.size}')
        check('exif' not in im.info and 'xmp' not in im.info, 'webp: no EXIF, no XMP')
        check(not re.findall(rb'(EXIF|XMP |C2PA)....', data[12:]) and not data[20] & 0x0C and int.from_bytes(data[4:8], 'little') == len(data) - 8, 'webp: no metadata chunks, flags cleared, RIFF size right')
        return
    if ext in ('odt', 'ods', 'odp'):
        import xml.dom.minidom
        z = zipfile.ZipFile(io.BytesIO(data))
        check(z.testzip() is None, f'{ext}: zip CRCs pass')
        first = z.infolist()[0]
        check(first.filename == 'mimetype' and first.compress_type == zipfile.ZIP_STORED, f'{ext}: mimetype stays first and stored')
        names = z.namelist()
        meta = z.read('meta.xml') + z.read('settings.xml') + z.read('META-INF/manifest.xml')
        leaks = [l for l in LEAKS + [b'Priya', b'Delacroix', b'HP '] if l in meta]
        check(not leaks, f'{ext}: no leaked strings in meta, settings or manifest ({leaks})')
        check(not [n for n in names if n.startswith('Thumbnails/')] and b'Thumbnails' not in z.read('META-INF/manifest.xml'), f'{ext}: thumbnail and its manifest entry gone')
        orig = zipfile.ZipFile(io.BytesIO(before))
        # Comments and the change log go (changes accepted); every other byte of content.xml stays.
        review = rb'<(office:annotation|text:tracked-changes)\b[^>]*>.*?<\/\1>|<(office:annotation-end|text:change|text:change-start|text:change-end)\b[^>]*\/>'
        check(z.read('content.xml') == re.sub(review, b'', orig.read('content.xml'), flags=re.S) and b'Delacroix' not in z.read('content.xml'), f'{ext}: content.xml byte for byte, less comments and the change log')
        for n in ('meta.xml', 'settings.xml', 'META-INF/manifest.xml'):
            xml.dom.minidom.parseString(z.read(n))
        return
    if ext == 'heic':
        import pillow_heif
        check(not leaks, f'heic: no leaked strings ({leaks})')
        check(len(data) == len(before), 'heic: same length, blanked in place')
        h = pillow_heif.open_heif(io.BytesIO(data))
        from PIL import Image
        orig = pillow_heif.open_heif(io.BytesIO(before))
        check(h.size == orig.size and bytes(h.data) == bytes(orig.data), f'heic: decodes, picture identical {h.size}')
        ex = Image.Exif(); ex.load(h.info.get('exif') or b'')
        check(not dict(ex) and not ex.get_ifd(0x8825), f'heic: EXIF item has no tags ({dict(ex)})')
        check(b'Villanueva' not in (h.info.get('xmp') or b''), 'heic: XMP item empty')
        return
    if ext in ('jpg', 'png'):
        check(not leaks, f'{ext}: no leaked strings ({leaks})')
        from PIL import Image
        im = Image.open(io.BytesIO(data))
        im.load()
        check(im.size[0] > 0, f'{ext}: Pillow opens it, {im.size}')
        check(not dict(im.getexif()), f'{ext}: no EXIF')
        check('xmp' not in im.info and 'comment' not in im.info, f'{ext}: no XMP, no comment')
        if ext == 'jpg':
            check(all(a[0] not in ('APP1', 'APP11', 'APP13', 'COM') for a in im.applist), f'{ext}: only {[a[0] for a in im.applist]} segments left')
        else:
            check(not im.text and 'exif' not in im.info, f'{ext}: no text chunks, no eXIf ({im.text})')
            chunks = re.findall(rb'(?s)....(tEXt|zTXt|iTXt|eXIf|tIME|caBX)', data)
            check(not chunks, f'{ext}: no metadata chunk types in the bytes ({chunks})')
        return
    z = zipfile.ZipFile(io.BytesIO(data))
    check(z.testzip() is None, f'{ext}: zip CRCs pass')
    names = z.namelist()
    # The document's own text may name people (a letterhead, a signature); metadata lives in docProps.
    props = b''.join(z.read(n) for n in names if n.startswith('docProps/'))
    leaks = [l for l in LEAKS if l in props]
    check(not leaks, f'{ext}: no leaked strings in docProps ({leaks})')
    core = z.read('docProps/core.xml')
    check(b'<dc:creator' not in core and b'lastModifiedBy' not in core and b'revision' not in core, f'{ext}: core.xml empty')
    app = z.read('docProps/app.xml')
    check(b'Company' not in app and b'TotalTime' not in app and b'Application' not in app, f'{ext}: app.xml empty')
    check(b'ClientMatter' not in z.read('docProps/custom.xml'), f'{ext}: custom.xml empty')
    check(not [n for n in names if 'thumbnail' in n.lower()], f'{ext}: no thumbnail part')
    check(all(i.date_time[:3] == (1980, 1, 1) for i in z.infolist()), f'{ext}: zip entry timestamps reset')
    if ext == 'docx':
        import docx
        d = docx.Document(io.BytesIO(data))
        check('Offer of employment' in d.paragraphs[0].text, 'docx: python-docx opens it, body intact')
        check(not d.core_properties.author and not d.core_properties.last_modified_by, 'docx: core properties empty')
    elif ext == 'xlsx':
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data))
        check(wb.active['B2'].value == 4200, 'xlsx: openpyxl opens it, cell intact')
        check(wb.properties.creator in (None, '', 'openpyxl'), 'xlsx: creator empty (openpyxl fills its own name when the field is absent)')
    else:
        from pptx import Presentation
        p = Presentation(io.BytesIO(data))
        check(p.slides[0].shapes.title.text == 'Pitch, internal only', 'pptx: python-pptx opens it, title intact')
        check(not p.core_properties.author, 'pptx: author empty')



def rel_target(rels, target):
    """A relationship's target as a part name."""
    if target.startswith('/'):
        return target[1:]
    parts = [p for p in re.sub(r'(^|/)_rels/[^/]+$', '', rels).split('/') if p]
    for s in target.split('/'):
        if s == '..':
            parts.pop()
        elif s and s != '.':
            parts.append(s)
    return '/'.join(parts)


def sound(name, data):
    """The rewritten package opens: CRCs pass, every XML part parses, and every relationship, content-type override and
    r:id points at something that exists."""
    import xml.dom.minidom
    z = zipfile.ZipFile(io.BytesIO(data))
    check(z.testzip() is None, f'{name}: zip CRCs and central directory pass')
    names = set(z.namelist())
    bad, dangling = [], []
    for n in sorted(names):
        if not n.endswith(('.xml', '.rels', '.vml')):
            continue
        try:
            xml.dom.minidom.parseString(z.read(n))
        except Exception as e:
            bad.append(f'{n}: {e}')
        if n.endswith('.rels'):
            rels = z.read(n).decode()
            for r in re.findall(r'<Relationship\b[^>]*>', rels):
                t = re.search(r'Target="([^"]*)"', r).group(1)
                if 'External' not in r and rel_target(n, t) not in names:
                    dangling.append(f'{n} -> {t}')
            owner = re.sub(r'(^|/)_rels/([^/]+)\.rels$', r'\1\2', n)
            if owner in names:
                ids = set(re.findall(r'\bId="([^"]*)"', rels))
                dangling += [f'{owner} r:id={i}' for i in re.findall(r'\sr:id="([^"]*)"', z.read(owner).decode()) if i not in ids]
    if '[Content_Types].xml' in names:
        dangling += [p for p in re.findall(r'PartName="/([^"]*)"', z.read('[Content_Types].xml').decode()) if p not in names]
    check(not bad, f'{name}: every XML part is well-formed ({bad})')
    check(not dangling, f'{name}: no relationship, content type or r:id points at a removed part ({dangling})')
    return z


def review():
    print('comments and tracked changes: removed, changes accepted')
    files = ['review.docx', 'review.xlsx', 'review.pptx', 'review.odt', 'review.ods']
    rv = json.loads(subprocess.run(['node', os.path.join(HERE, 'lib.mjs'), *files, *[f'{n}@comments,changes' for n in files], 'review.docx@comments', 'review.docx@changes', 'review.xlsx@comments'], check=True, capture_output=True, text=True).stdout)
    got = lambda n: open(os.path.join(OUT, n), 'rb').read()
    GONE = [b'Okonkwo', b'Teodoro', b'Rask', b'mokonkwo', b'push this to June', b'Check with finance', b'Too optimistic', b'Second opinion', b'88,000', b'<v>4800', b'value="4800"']
    for name in files:
        r = rv[name]
        check('error' not in r and r.get('left') == [], f'{name}: read, and the clean copy read back has no comments or changes left ({r})')
        check('comments=comments' in r['keys'] and ('changes=changes' in r['keys']) != name.endswith('pptx'), f"{name}: comments{'' if name.endswith('pptx') else ' and changes'} are removable segments ({r['keys']})")
        check('Written by=Marguerite Okonkwo-Hale' in r['fields'], f'{name}: the comment author is named in the report')
        if not name.endswith('pptx'):
            check('Changed by=Teodoro Rask' in r['fields'], f'{name}: the author of the changes is named in the report')
        data = got(name)
        z = sound(name, data)
        inside = data + b''.join(z.read(n) for n in z.namelist())
        leaks = [g for g in GONE if g in inside]
        check(not leaks, f'{name}: comment text, deleted words and every comment or change author gone from all the bytes ({leaks})')
    import docx, openpyxl
    from pptx import Presentation
    z = zipfile.ZipFile(os.path.join(OUT, 'review.docx'))
    d = docx.Document(io.BytesIO(got('review.docx')))
    text = '\n'.join(p.text for p in d.paragraphs)
    check('The salary is 92,000 per year.' in text and 'Start date is May.' in text and 'Closing line.' in text and text.count('Moved sentence.') == 1, f'review.docx: python-docx opens it; insertion kept, deletion gone, moved text once ({text!r})')
    check([c.text for row in d.tables[0].rows for c in row.cells] == ['Row that stays'], 'review.docx: the deleted table row is gone, the other row kept')
    check(len(d.tables) == 1 and b'Whole table deleted' not in z.read('word/document.xml'), 'review.docx: a table whose rows were all deleted goes with them')
    check(any(r.bold and r.text == 'Bold now.' for p in d.paragraphs for r in p.runs), 'review.docx: the formatting change is accepted (bold stays, the old formatting record goes)')
    body = z.read('word/document.xml')
    check(not re.findall(rb'<w:(ins|del|moveFrom|moveTo|rPrChange|commentRangeStart|commentReference)\b', body) and b'trackRevisions' not in z.read('word/settings.xml'), 'review.docx: no change or comment marks left, track changes switched off')
    x = zipfile.ZipFile(os.path.join(OUT, 'review.xlsx'))
    wb = openpyxl.load_workbook(io.BytesIO(got('review.xlsx')))
    check(wb.active['B2'].value == 5100 and wb.active['B2'].comment is None, 'review.xlsx: openpyxl opens it, the value stays, the comment is gone')
    vml = x.read('xl/drawings/commentsDrawing1.vml')
    check(b'ObjectType="Button"' in vml and b'ObjectType="Note"' not in vml and b'legacyDrawing' in x.read('xl/worksheets/sheet1.xml'), 'review.xlsx: the comment box leaves the VML, the form button in it stays')
    check(not [n for n in x.namelist() if re.match(r'xl/(revisions|persons|threadedComments|comments)', n)], 'review.xlsx: revision log, people, threaded and legacy comments gone')
    p = Presentation(io.BytesIO(got('review.pptx')))
    check(p.slides[0].shapes.title.text == 'Quarterly plan', 'review.pptx: python-pptx opens it, title intact')
    check(b'commentRel' not in zipfile.ZipFile(os.path.join(OUT, 'review.pptx')).read('ppt/slides/slide1.xml'), 'review.pptx: the slide no longer links to modern comments')
    o = zipfile.ZipFile(os.path.join(OUT, 'review.odt'))
    c = o.read('content.xml')
    check(b'The salary is 92,000 per year.' in c and b'Start date is May.' in c and not re.findall(rb'<(text:tracked-changes|text:change|office:annotation)\b', c), 'review.odt: insertion kept, deletion, change marks and comment gone')
    s = zipfile.ZipFile(os.path.join(OUT, 'review.ods')).read('content.xml')
    check(b'office:value="5100"' in s and b'Kestrel Ltd' in s and b'tracked-changes' not in s and b'annotation' not in s, 'review.ods: the cell value stays, its comment and the change log go')
    check(b'>false<' in o.read('settings.xml') and o.infolist()[0].filename == 'mimetype' and o.infolist()[0].compress_type == zipfile.ZIP_STORED, 'review.odt: track changes off, mimetype first and stored')
    print('comments and tracked changes: kept by policy')
    for name in files:
        k = f'{name}@comments,changes'
        check(rv[k].get('left') == [] and rv[k]['removed'] > 0, f'{k}: the rest is still cleaned ({rv[k]})')
        a, b = zipfile.ZipFile(os.path.join(HERE, 'files', name)), zipfile.ZipFile(os.path.join(OUT, f'comments+changes-{name}'))
        same = [n for n in a.namelist() if not n.startswith(('docProps/', 'meta.xml', 'Thumbnails/', '_rels/.rels'))]
        check(all(n in b.namelist() and a.read(n) == b.read(n) for n in same), f'{k}: every content part byte for byte')
    z = sound('review.docx@comments', got('comments-review.docx'))
    check(b'Okonkwo' in z.read('word/comments.xml') and b'w:del ' not in z.read('word/document.xml') and b'commentReference' in z.read('word/document.xml'), 'review.docx@comments: comments kept, changes accepted')
    z = sound('review.docx@changes', got('changes-review.docx'))
    check('word/comments.xml' not in z.namelist() and b'w:del ' in z.read('word/document.xml') and b'commentReference' not in z.read('word/document.xml'), 'review.docx@changes: changes kept, comments gone')
    z = sound('review.xlsx@comments', got('comments-review.xlsx'))
    check('xl/comments/comment1.xml' in z.namelist() and not [n for n in z.namelist() if n.startswith('xl/revisions/')], 'review.xlsx@comments: comments kept, revision log gone')


if __name__ == '__main__':
    rep = json.loads(subprocess.run(['node', os.path.join(HERE, 'lib.mjs'), *CASES, 'notajpeg.jpg', 'notes.txt', 'clean.jpg'], check=True, capture_output=True, text=True).stdout)
    for name in CASES:
        print(name)
        check('error' not in rep[name], f'{name}: read ({rep[name]})')
        check(rep[name].get('removed', 0) > 0, f'{name}: fields found to remove')
        check(rep[name].get('left') == [], f'{name}: the clean copy, read back, has no metadata left ({rep[name].get("left")})')
        before = open(os.path.join(HERE, 'files', name), 'rb').read()
        verify(name, open(os.path.join(OUT, name), 'rb').read(), before, rep[name])
    print('c2pa')
    for name in ('c2pa.jpg', 'c2pa.png'):
        f = dict(x.split('=', 1) for x in rep[name]['fields'])
        check(f.get('Made with (Content Credentials)') == 'make_test_images 0.33.1', f'{name}: claim generator read ({f})')
        check(f.get('Signed by') == 'C2PA Test Signing Cert', f'{name}: signer read')
        check('opened' in f.get('Edit history', '') and f.get('Made from') == '1 earlier file', f'{name}: actions and ingredient read')
    bz = subprocess.run(['node', os.path.join(HERE, 'batch.mjs')], capture_output=True, text=True)
    check(bz.returncode == 0, f'batch: a zip opens into its files with folders kept, and a record carries fingerprints and the read-back ({bz.stderr.strip()[-300:]})')
    nm = subprocess.run(['node', os.path.join(HERE, 'name.mjs')], capture_output=True, text=True)
    check(nm.returncode == 0, f'name: dates, emails and draft words in a file name are flagged ({nm.stderr.strip()[-300:]})')
    ai = subprocess.run(['node', os.path.join(HERE, 'c2pa.mjs')], capture_output=True, text=True)
    check(ai.returncode == 0, f'c2pa: an AI generator and source type are named ({ai.stderr.strip()[-200:]})')
    print('content checks: shown, kept')
    risk = json.loads(subprocess.run(['node', os.path.join(HERE, 'lib.mjs'), 'hidden.xlsx', 'hidden.docx', 'hidden.pptx', 'saves.pdf'], check=True, capture_output=True, text=True).stdout)
    want = {
        'hidden.xlsx': ['Hidden sheet=Salaries', 'Very hidden sheet=Owners', 'Hidden rows and columns=Fees: 1 row, 1 column', 'Pivot table data=943 source rows', 'Comments=1 comment', 'Written by=H. Delacroix', 'Linked workbook=file:///C:/Users/hdelacroix/Clients/Tamarind rates.xlsx'],
        'hidden.docx': ['Hidden text=1 passage', 'Deleted text=1 tracked deletion', 'Changed by=H. Delacroix', 'Fetches on opening=https://canary.example/t.png', 'Image fetched on opening=https://pixel.example/open?id=7'],
        'hidden.pptx': ['Speaker notes=on 1 slide', 'Hidden slides=1 slide', 'Embedded file=Microsoft_Excel_Worksheet.xlsx'],
        'saves.pdf': ['Earlier versions inside=1:', 'Runs JavaScript=', 'Contacts on opening=https://track.example/o?doc=31'],
    }
    for name, fields in want.items():
        got = risk[name].get('fields', [])
        for w in fields:
            check(any(g.startswith(w) for g in got), f'{name}: reports {w}')
        check('filesanity.com' not in ' '.join(got), f'{name}: an ordinary hyperlink is not reported') if name == 'hidden.docx' else None
        out = open(os.path.join(OUT, name), 'rb').read()
        if name.endswith('pdf'):
            check(out.count(b'startxref') == 2 and b'/OpenAction' in out and len(out) == os.path.getsize(os.path.join(HERE, 'files', name)), 'saves.pdf: cleaning leaves the saves and actions as they are')
        else:
            z = zipfile.ZipFile(io.BytesIO(out))
            check(z.testzip() is None, f'{name}: cleaned copy is a sound zip')
    z = zipfile.ZipFile(os.path.join(OUT, 'hidden.xlsx'))
    check(b'state="hidden"' in z.read('xl/workbook.xml') and 'xl/pivotCache/pivotCacheRecords1.xml' in z.namelist(), 'hidden.xlsx: the hidden sheet and pivot cache are left in the file')
    check(not [n for n in z.namelist() if 'comment' in n.lower()] and b'Delacroix' not in z.read('xl/worksheets/sheet1.xml'), 'hidden.xlsx: the cell comment and its box are gone')
    doc = zipfile.ZipFile(os.path.join(OUT, 'hidden.docx')).read('word/document.xml')
    check(b'w:delText' not in doc and b'Salary floor' not in doc and b'Salary as agreed.' in doc and b'Delacroix' not in doc, 'hidden.docx: the tracked deletion is gone, the insertion kept as plain text')
    review()
    lk = json.loads(subprocess.run(['node', os.path.join(HERE, 'lib.mjs'), 'locked.pdf'], check=True, capture_output=True, text=True).stdout)['locked.pdf']
    check(lk['removed'] == 0 and lk['kept'] == 1 and 'encrypted' in lk['note'] and 'not cleaned' in lk['note'], f'locked.pdf: an encrypted PDF is shown as not cleaned, nothing blanked ({lk})')
    check(open(os.path.join(OUT, 'locked.pdf'), 'rb').read() == open(os.path.join(HERE, 'files', 'locked.pdf'), 'rb').read(), 'locked.pdf: left byte for byte')
    print('edge cases')
    EDGE = ['rotated.jpg', 'fill.jpg', 'motion.jpg', 'multi.jpg', 'info-ref.pdf', 'nested.pdf', 'quote-xmp.pdf', 'bare-xmp.pdf', 'locked-lin.pdf', 'memo-z.pdf']
    ed = json.loads(subprocess.run(['node', os.path.join(HERE, 'lib.mjs'), *EDGE], check=True, capture_output=True, text=True).stdout)
    got = lambda n: open(os.path.join(OUT, n), 'rb').read()
    for name in EDGE[:-2]:
        check(ed[name].get('left') == [], f"{name}: read back, nothing left ({ed[name].get('left', ed[name])})")
    from PIL import Image
    ex = Image.open(io.BytesIO(got('rotated.jpg'))).getexif()
    check(dict(ex) == {0x0112: 6} and not [l for l in LEAKS if l in got('rotated.jpg')], f'rotated.jpg: Orientation=6 kept, Artist and GPS gone ({dict(ex)})')
    check('Artist=Jane Secret' in ed['rotated.jpg']['fields'] and ed['rotated.jpg']['removed'] > 1, 'rotated.jpg: Artist and GPS reported')
    for name in ('fill.jpg', 'motion.jpg', 'multi.jpg'):
        data = got(name)
        im = Image.open(io.BytesIO(data))
        im.load()
        leaks = [l for l in LEAKS + [b'ftyp', b'MPF'] if l in data]
        check(not leaks and not dict(im.getexif()) and data.endswith(b'\xff\xd9') and data.count(b'\xff\xd9') == 1, f'{name}: opens, one image, no EXIF, nothing after EOI, no leaks ({leaks})')
    check('Artist=Jane Secret' in ed['fill.jpg']['fields'] and any(f.startswith('GPS latitude=') for f in ed['fill.jpg']['fields']), f"fill.jpg: EXIF behind fill bytes reported ({ed['fill.jpg']['fields']})")
    for name in ('motion.jpg', 'multi.jpg'):
        check(any(f.startswith('Data after the image=') for f in ed[name]['fields']), f"{name}: data after the image reported as removed ({ed[name]['fields']})")
    want = {'info-ref.pdf': ['Author=Jane Secret', 'Title=Jane Secret CV'], 'nested.pdf': ['Title=Report (draft (v2)) for Jane Secret', 'Author=Ann ) Byrne-Ortiz \\'],
            'quote-xmp.pdf': ["pdf:Author=Sean O'Brien"], 'bare-xmp.pdf': ['pdf:Author=Jane Secret']}
    for name, fields in want.items():
        print(name)
        for w in fields:
            check(w in ed[name]['fields'], f"{name}: reports {w} ({ed[name]['fields']})")
        verify(name, got(name), open(os.path.join(HERE, 'files', name), 'rb').read(), ed[name])
    ll = ed['locked-lin.pdf']
    check(ll['removed'] == 0 and ll['kept'] == 1 and 'not cleaned' in ll['note'] and got('locked-lin.pdf') == open(os.path.join(HERE, 'files', 'locked-lin.pdf'), 'rb').read(), f'locked-lin.pdf: a linearized encrypted PDF is shown as not cleaned, left byte for byte ({ll})')
    check(any(k.startswith('xmpz-') and k.endswith('=xmp') for k in ed['memo-z.pdf']['keys']), f"memo-z.pdf: compressed XMP answers to the xmp policy key ({ed['memo-z.pdf']['keys']})")
    print('photos inside documents')
    pi = json.loads(subprocess.run(['node', os.path.join(HERE, 'lib.mjs'), 'photo-in.docx', 'photo-in.odt', 'photo-in.pdf'], check=True, capture_output=True, text=True).stdout)
    for name in pi:
        check(pi[name].get('left') == [] and any('Artist=Jane Secret' in f for f in pi[name].get('fields', [])), f"{name}: the photo's Artist is reported, nothing left on read-back ({pi[name].get('left', pi[name])})")
    for name in ('photo-in.docx', 'photo-in.odt'):
        z = zipfile.ZipFile(os.path.join(OUT, name))
        check(z.testzip() is None, f'{name}: cleaned copy is a sound zip')
        for n in [n for n in z.namelist() if re.search(r'(media|Pictures)/.*\.(jpe?g|png)$', n)]:
            d = z.read(n)
            im = Image.open(io.BytesIO(d))
            im.load()
            check(not dict(im.getexif()) and not im.text if hasattr(im, 'text') else not dict(im.getexif()), f'{name}: {n} opens with no EXIF or text')
            check(b'Jane Secret' not in d, f'{name}: {n} has no leaked name')
    data = got('photo-in.pdf')
    check(len(data) == os.path.getsize(os.path.join(HERE, 'files', 'photo-in.pdf')) and b'Jane Secret' not in data, 'photo-in.pdf: same length, the photo\'s EXIF blanked')
    import pikepdf
    with pikepdf.open(os.path.join(OUT, 'photo-in.pdf')) as p:
        Image.open(io.BytesIO(p.get_object(8, 0).read_raw_bytes())).load()
        check(not p.check_pdf_syntax() if hasattr(p, 'check_pdf_syntax') else True, 'photo-in.pdf: qpdf finds no syntax problems, the photo still decodes')
    print('refusals')
    check('does not begin like a JPEG' in rep['notajpeg.jpg'].get('error', ''), 'a .jpg that is not a JPEG is refused by its bytes')
    check('cannot read .txt' in rep['notes.txt'].get('error', ''), '.txt is refused in plain words')
    check(rep['clean.jpg'].get('removed') == 0, 'an already clean JPEG reports nothing to remove')
    print(f"\n{len(failures)} failures" if failures else '\nall passed')
    sys.exit(1 if failures else 0)
