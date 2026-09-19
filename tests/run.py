"""Clean every file in tests/files through the built page's picker and assert
the metadata is gone and the file still opens. Also: the picker's states,
a synthetic drop, paste, focus after a keyboard clean, and zero network
requests during read and clean in Chromium and Firefox.

Needs the preview server: npx vite preview --port 4184
Run: python3 tests/run.py [url]
"""
import io
import json
import os
import re
import sys
import zipfile

from playwright.sync_api import sync_playwright

URL = next((a for a in sys.argv[1:] if a.startswith('http')), 'http://127.0.0.1:4184/')
HERE = os.path.dirname(os.path.abspath(__file__))
FILES = os.path.join(HERE, 'files')
OUT = os.path.join(HERE, 'out')
os.makedirs(OUT, exist_ok=True)

CASES = ['photo.jpg', 'shot.png', 'offer.docx', 'fees.xlsx', 'pitch.pptx', 'memo.pdf', 'memo-z.pdf']
LEAKS = [b'Okafor', b'Priya', b'Delacroix', b'Lightroom', b'GIMP', b'iPhone', b'Oxford', b'Northwind', b'DEL-2026', b'Byrne', b'confidential']

failures = []


def check(cond, msg):
    print(('  ok   ' if cond else '  FAIL ') + msg)
    if not cond:
        failures.append(msg)


def verify(name, data, before, fig=''):
    ext = name.rsplit('.', 1)[1]
    leaks = [l for l in LEAKS if l in data]
    if ext == 'pdf':
        # Same length, blanked in place: the words must be gone but the structure untouched.
        check(len(data) == len(before), 'pdf: same length as the original')
        check(not leaks, f'pdf: no leaked strings ({leaks})')
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(data))
        check(len(r.pages) == 1 and 'Internal memo' in r.pages[0].extract_text(), 'pdf: opens and the page text is intact')
        info = {k: str(v).strip() for k, v in (r.metadata or {}).items()}
        check(all(not v for v in info.values()), f'pdf: Info values blank ({info})')
        xmp = r.xmp_metadata
        blank = lambda v: not ''.join(v if isinstance(v, list) else [v or '']).strip()
        if name == 'memo-z.pdf':
            # The compressed stream is kept, and the page must say so rather than claim a clean file.
            check(xmp is not None and not blank(xmp.dc_creator), f'pdf-z: compressed XMP kept, creator still present ({xmp and xmp.dc_creator})')
            check('compressed XMP stream' in fig and 'shown and kept' in fig, 'pdf-z: the note says the compressed stream is shown and kept')
            check('8 removed' in fig and '1 field' in fig.split('After')[-1], 'pdf-z: the after figure counts the kept field')
        else:
            check(xmp is None or (blank(xmp.dc_creator) and blank(xmp.xmp_creator_tool)), f'pdf: XMP creator and tool blank ({xmp and xmp.dc_creator}, {xmp and xmp.xmp_creator_tool})')
            check('0 fields' in fig.split('After')[-1], 'pdf: the after figure reads 0 fields')
        return
    check(not leaks, f'{ext}: no leaked strings ({leaks})')
    if ext in ('jpg', 'png'):
        from PIL import Image
        im = Image.open(io.BytesIO(data))
        im.load()
        check(im.size[0] > 0, f'{ext}: Pillow opens it, {im.size}')
        check(not dict(im.getexif()), f'{ext}: no EXIF')
        check('xmp' not in im.info and 'comment' not in im.info, f'{ext}: no XMP, no comment')
        if ext == 'jpg':
            check(all(a[0] not in ('APP1', 'APP13', 'COM') for a in im.applist), f'{ext}: only {[a[0] for a in im.applist]} segments left')
        else:
            check(not im.text and 'exif' not in im.info, f'{ext}: no text chunks, no eXIf ({im.text})')
            chunks = re.findall(rb'(?s)....(tEXt|zTXt|iTXt|eXIf|tIME)', data)
            check(not chunks, f'{ext}: no metadata chunk types in the bytes ({chunks})')
        return
    z = zipfile.ZipFile(io.BytesIO(data))
    check(z.testzip() is None, f'{ext}: zip CRCs pass')
    names = z.namelist()
    core = z.read('docProps/core.xml')
    check(b'<dc:creator' not in core and b'lastModifiedBy' not in core and b'revision' not in core, f'{ext}: core.xml empty')
    app = z.read('docProps/app.xml')
    check(b'Company' not in app and b'TotalTime' not in app and b'Application' not in app, f'{ext}: app.xml empty')
    check(b'ClientMatter' not in z.read('docProps/custom.xml'), f'{ext}: custom.xml empty')
    check(not [n for n in names if 'thumbnail' in n.lower()], f'{ext}: no thumbnail part')
    check(all(i.date_time[:3] == (1980, 1, 1) for i in z.infolist()), f'{ext}: zip entry timestamps reset')
    if ext == 'docx':
        import docx
        d = docx.Document(io.BytesIO(data))
        check('Offer letter draft' in d.paragraphs[0].text, 'docx: python-docx opens it, body intact')
        check(not d.core_properties.author and not d.core_properties.last_modified_by, 'docx: core properties empty')
    elif ext == 'xlsx':
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data))
        check(wb.active['B2'].value == 4200, 'xlsx: openpyxl opens it, cell intact')
        check(wb.properties.creator in (None, '', 'openpyxl'), 'xlsx: creator empty (openpyxl fills its own name when the field is absent)')
    else:
        from pptx import Presentation
        p = Presentation(io.BytesIO(data))
        check(p.slides[0].shapes.title.text == 'Pitch, internal only', 'pptx: python-pptx opens it, title intact')
        check(not p.core_properties.author, 'pptx: author empty')


PICKER = '[data-state]'
HERO = 'section[aria-labelledby=h1]'


def run_cases(page, requests):
    n0 = len(requests)
    for name in CASES:
        print(name)
        path = os.path.join(FILES, name)
        page.set_input_files('input[type=file]', path)
        page.wait_for_selector('button:has-text("Remove all")', timeout=8000)
        check(page.get_attribute(PICKER, 'data-state') == 'ready', 'state is ready')
        page.focus('button:has-text("Remove all")')
        with page.expect_download() as dl:
            page.keyboard.press('Enter')
        out = os.path.join(OUT, dl.value.suggested_filename)
        dl.value.save_as(out)
        page.wait_for_selector('button:has-text("Download again")', timeout=8000)
        with open(out, 'rb') as f:
            data = f.read()
        with open(path, 'rb') as f:
            before = f.read()
        check(dl.value.suggested_filename == name.replace('.', '-clean.'), f'named {dl.value.suggested_filename}')
        page.wait_for_timeout(150)
        active = page.evaluate('document.activeElement.textContent')
        check(active == 'Download again', f'focus after keyboard clean is on the download button ({active!r})')
        live = page.inner_text('[role=status]')
        check('cleaned' in live and 'removed' in live, f'the live region announces the result ({live!r})')
        check(page.get_attribute(PICKER, 'data-state') == 'cleaned', 'state is cleaned')
        verify(name, data, before, page.inner_text(PICKER))
    return n0


with sync_playwright() as p:
    b = p.chromium.launch()
    ctx = b.new_context(accept_downloads=True)
    page = ctx.new_page()
    errors = []
    page.on('pageerror', lambda e: errors.append(str(e)))
    requests = []
    # blob: URLs are the page's own object URLs (download); they never leave the browser.
    page.on('request', lambda r: r.url.startswith('blob:') or requests.append(r.url))
    page.goto(URL, wait_until='load')
    page.wait_for_timeout(500)
    # The native input is never rendered: no "Choose file" control is visible, and the surface is a real button.
    check(page.evaluate("() => { const i = document.querySelector('input[type=file]'); const r = i.getBoundingClientRect(); return r.width <= 1 && r.height <= 1 && i.getAttribute('aria-hidden') === 'true' && i.tabIndex === -1 }"), 'the native file input has no pixels and is out of the tab order')
    check(page.get_attribute(PICKER, 'data-state') == 'idle', 'state is idle on load')
    surface = page.locator('button.drop')
    check(surface.count() == 1, 'the drop surface is a button')
    # Enter and Space on the surface open the picker: the OS dialog cannot be observed, so the input's click is.
    page.evaluate("() => { window.__clicks = 0; HTMLInputElement.prototype.click = function () { if (this.type === 'file') window.__clicks++ } }")
    surface.focus()
    page.keyboard.press('Enter')
    page.keyboard.press('Space')
    check(page.evaluate('window.__clicks') == 2, 'Enter and Space on the surface open the picker')
    ring = surface.evaluate('e => { e.focus(); const s = getComputedStyle(e); return s.outlineStyle + " " + s.outlineWidth }')
    check(ring.startswith('solid'), f'the surface has a visible focus ring ({ring})')
    # Drag over and a synthetic drop of a real file.
    dt = page.evaluate_handle("async () => { const r = await fetch('/sample.jpg'); const b = await r.blob(); const dt = new DataTransfer(); dt.items.add(new File([b], 'dropped.jpg', { type: 'image/jpeg' })); return dt }")
    page.dispatch_event(HERO, 'dragenter', {'dataTransfer': dt})
    page.wait_for_timeout(100)
    check('Release to read it' in page.inner_text(PICKER), 'drag-over state shows on the surface')
    page.dispatch_event(HERO, 'drop', {'dataTransfer': dt})
    page.wait_for_selector('button:has-text("Remove all")', timeout=8000)
    check('dropped.jpg' in page.inner_text(PICKER), 'a synthetic drop reads the file')
    page.click('button[aria-label="Remove dropped.jpg"]')
    page.wait_for_timeout(200)
    check(page.get_attribute(PICKER, 'data-state') == 'idle', "the chip's remove control returns to idle")
    requests.clear()
    n0 = run_cases(page, requests)
    # Edge states.
    page.set_input_files('input[type=file]', os.path.join(FILES, 'clean.jpg'))
    page.wait_for_timeout(600)
    check('Nothing to remove' in page.inner_text(PICKER), 'clean.jpg: nothing to remove')
    page.set_input_files('input[type=file]', os.path.join(FILES, 'notes.txt'))
    page.wait_for_timeout(600)
    check(page.get_attribute(PICKER, 'data-state') == 'error' and 'cannot read .txt files yet' in page.inner_text(PICKER), 'notes.txt: error state, unsupported by kind')
    page.set_input_files('input[type=file]', os.path.join(FILES, 'notajpeg.jpg'))
    page.wait_for_timeout(600)
    check('This .jpg does not begin like a JPEG' in page.inner_text(PICKER), 'notajpeg.jpg: refused by first bytes')
    # The picker is single-file; two files only arrive by drop, so drop them.
    dt = page.evaluate_handle("() => { const dt = new DataTransfer(); dt.items.add(new File([new Uint8Array([0xff, 0xd8, 0xff, 0xe0])], 'photo.jpg')); dt.items.add(new File(['x'], 'shot.png')); return dt }")
    page.dispatch_event(HERO, 'drop', {'dataTransfer': dt})
    page.wait_for_timeout(600)
    check('One file at a time. Reading photo.jpg.' in page.inner_text(PICKER), 'two files: reads the first and says so')
    # Paste.
    page.evaluate("() => { const dt = new DataTransfer(); dt.items.add(new File([new Uint8Array([0x89, 0x50, 0x4e, 0x47, 0x0d, 0x0a, 0x1a, 0x0a])], 'pasted.png', { type: 'image/png' })); document.dispatchEvent(new ClipboardEvent('paste', { clipboardData: dt, bubbles: true })) }")
    page.wait_for_timeout(600)
    check('pasted.png' in page.inner_text(PICKER), 'a pasted file is read')
    check(not errors, f'no page errors ({errors})')
    check(len(requests) == n0, f'chromium: no network requests during {len(CASES)} reads and cleans ({requests[n0:]})')
    b.close()

    # Firefox: the client's browser. One read and clean, zero requests.
    b = p.firefox.launch()
    ctx = b.new_context(accept_downloads=True)
    page = ctx.new_page()
    requests = []
    page.on('request', lambda r: r.url.startswith('blob:') or requests.append(r.url))
    page.goto(URL, wait_until='load')
    page.wait_for_timeout(500)
    n0 = len(requests)
    page.set_input_files('input[type=file]', os.path.join(FILES, 'photo.jpg'))
    page.wait_for_selector('button:has-text("Remove all")', timeout=8000)
    with page.expect_download() as dl:
        page.click('button:has-text("Remove all")')
    page.wait_for_selector('button:has-text("Download again")', timeout=8000)
    check('0 fields' in page.inner_text(PICKER).split('After')[-1], 'firefox: the after figure reads 0 fields')
    check(len(requests) == n0, f'firefox: no network requests during a read and clean ({requests[n0:]})')
    b.close()

json.dump({'failures': failures}, open(os.path.join(OUT, 'result.json'), 'w'))
print('\nFAILED' if failures else '\nALL PASS', len(failures))
sys.exit(1 if failures else 0)
