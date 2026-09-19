"""Build the bundled sample JPEG and the test corpus.

public/sample.jpg   the Codex still, resized, with EXIF (camera, dates, GPS,
                    artist, software), XMP and IPTC injected
tests/files/*       one file per format the parser reads, each carrying
                    metadata, plus one clean JPEG and one unsupported file

Run: python3 tools/make_samples.py
"""
import io
import os
import struct
import zipfile

from PIL import Image
from PIL.TiffImagePlugin import IFDRational as R

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'assets-src', 'sample-src.png')
PUB = os.path.join(ROOT, 'public')
TST = os.path.join(ROOT, 'tests', 'files')
os.makedirs(PUB, exist_ok=True)
os.makedirs(TST, exist_ok=True)

# All values are mock: a plausible phone photo. Marked here, not on the page.
XMP = b'''<?xpacket begin="\xef\xbb\xbf" id="W5M0MpCehiHzreSzNTczkc9d"?>
<x:xmpmeta xmlns:x="adobe:ns:meta/" x:xmptk="Adobe XMP Core 9.1-c001">
 <rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#">
  <rdf:Description rdf:about=""
    xmlns:xmp="http://ns.adobe.com/xap/1.0/"
    xmlns:dc="http://purl.org/dc/elements/1.1/"
    xmlns:photoshop="http://ns.adobe.com/photoshop/1.0/"
    xmlns:xmpMM="http://ns.adobe.com/xap/1.0/mm/"
    xmlns:stEvt="http://ns.adobe.com/xap/1.0/sType/ResourceEvent#"
    xmp:CreatorTool="Adobe Lightroom 8.2 (iOS)"
    xmp:CreateDate="2026-04-11T17:42:09+01:00"
    xmp:ModifyDate="2026-04-12T09:03:51+01:00"
    photoshop:City="Oxford"
    photoshop:Country="United Kingdom"
    xmpMM:DocumentID="xmp.did:8f2e1a40-6c2b-4b2e-9c53-2a5f0f2d9b11"
    xmpMM:OriginalDocumentID="xmp.did:8f2e1a40-6c2b-4b2e-9c53-2a5f0f2d9b11">
   <dc:creator><rdf:Seq><rdf:li>M. Okafor</rdf:li></rdf:Seq></dc:creator>
   <dc:rights><rdf:Alt><rdf:li xml:lang="x-default">M. Okafor, all rights reserved</rdf:li></rdf:Alt></dc:rights>
   <dc:subject><rdf:Bag><rdf:li>bicycle</rdf:li><rdf:li>oxford</rdf:li><rdf:li>for sale</rdf:li></rdf:Bag></dc:subject>
   <xmpMM:History><rdf:Seq>
    <rdf:li stEvt:action="created" stEvt:when="2026-04-11T17:42:09+01:00" stEvt:softwareAgent="iOS 19.4"/>
    <rdf:li stEvt:action="saved" stEvt:when="2026-04-12T09:03:51+01:00" stEvt:softwareAgent="Adobe Lightroom 8.2 (iOS)"/>
   </rdf:Seq></xmpMM:History>
  </rdf:Description>
 </rdf:RDF>
</x:xmpmeta>
<?xpacket end="w"?>'''


def exif_bytes():
    ex = Image.Exif()
    ex[0x010F] = 'Apple'
    ex[0x0110] = 'iPhone 16 Pro'
    ex[0x0131] = 'Adobe Lightroom 8.2 (iOS)'
    ex[0x0132] = '2026:04:12 09:03:51'
    ex[0x013B] = 'M. Okafor'
    ex[0x8298] = 'M. Okafor, all rights reserved'
    ex[0x0112] = 1
    sub = ex.get_ifd(0x8769)
    sub[0x9003] = '2026:04:11 17:42:09'
    sub[0x9004] = '2026:04:11 17:42:09'
    sub[0xA434] = 'iPhone 16 Pro back camera 6.765mm f/1.78'
    sub[0xA433] = 'Apple'
    sub[0x829A] = R(1, 240)
    sub[0x829D] = R(178, 100)
    sub[0x8827] = 64
    sub[0x920A] = R(6765, 1000)
    sub[0xA430] = 'Okafor'
    sub[0xA431] = 'F2LXK9Q0PLJ7'
    gps = ex.get_ifd(0x8825)
    gps[1] = 'N'
    gps[2] = (R(51, 1), R(45, 1), R(1932, 100))
    gps[3] = 'W'
    gps[4] = (R(1, 1), R(15, 1), R(4128, 100))
    gps[5] = 0
    gps[6] = R(61, 1)
    gps[7] = (R(16, 1), R(42, 1), R(9, 1))
    gps[29] = '2026:04:11'
    return ex.tobytes()


def iptc_app13():
    def ds(rec, tag, val):
        v = val.encode('utf-8')
        return bytes([0x1C, rec, tag]) + struct.pack('>H', len(v)) + v
    iptc = b''.join([
        ds(2, 5, 'Bicycle, Holywell Street'),
        ds(2, 25, 'bicycle'), ds(2, 25, 'oxford'), ds(2, 25, 'for sale'),
        ds(2, 55, '20260411'),
        ds(2, 80, 'M. Okafor'),
        ds(2, 90, 'Oxford'), ds(2, 101, 'United Kingdom'),
        ds(2, 116, 'M. Okafor, all rights reserved'),
        ds(2, 120, 'Black city bicycle against a brick wall, listed for sale.'),
    ])
    irb = b'8BIM' + struct.pack('>H', 0x0404) + b'\x00\x00' + struct.pack('>I', len(iptc)) + iptc
    if len(iptc) % 2:
        irb += b'\x00'
    payload = b'Photoshop 3.0\x00' + irb
    return b'\xFF\xED' + struct.pack('>H', len(payload) + 2) + payload


def insert_after_app_segments(jpeg, segment):
    # Insert a new segment after the last APPn segment in the header.
    i = 2
    while i < len(jpeg):
        marker = jpeg[i + 1]
        if not (0xE0 <= marker <= 0xEF):
            break
        i += 2 + struct.unpack('>H', jpeg[i + 2:i + 4])[0]
    return jpeg[:i] + segment + jpeg[i:]


def main():
    im = Image.open(SRC).convert('RGB')
    im = im.resize((1200, 800), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, 'JPEG', quality=72, exif=exif_bytes(), xmp=XMP, comment=b'Sent from my phone')
    jpeg = insert_after_app_segments(buf.getvalue(), iptc_app13())
    with open(os.path.join(PUB, 'sample.jpg'), 'wb') as f:
        f.write(jpeg)
    with open(os.path.join(TST, 'photo.jpg'), 'wb') as f:
        f.write(jpeg)
    print('sample.jpg', len(jpeg), 'bytes')

    # A JPEG with nothing to remove.
    clean = io.BytesIO()
    im.resize((300, 200)).save(clean, 'JPEG', quality=70)
    open(os.path.join(TST, 'clean.jpg'), 'wb').write(clean.getvalue())

    # PNG with tEXt, zTXt, iTXt, tIME and eXIf.
    from PIL import PngImagePlugin
    small = im.resize((320, 213))
    info = PngImagePlugin.PngInfo()
    info.add_text('Author', 'M. Okafor')
    info.add_text('Software', 'GIMP 3.0.4')
    info.add_text('Comment', 'exported for the listing', zip=True)
    info.add_itxt('Title', 'Vélo, Holywell Street', lang='fr', tkey='Titre')
    info.add(b'tIME', struct.pack('>HBBBBB', 2026, 4, 12, 9, 3, 51))
    info.add(b'eXIf', exif_bytes())
    small.save(os.path.join(TST, 'shot.png'), 'PNG', pnginfo=info)

    # DOCX, XLSX, PPTX with core, app and custom properties.
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    d = docx.Document()
    d.add_paragraph('Offer letter draft. Salary and start date to be confirmed.')
    cp = d.core_properties
    cp.author = 'Priya Raman'
    cp.last_modified_by = 'H. Delacroix'
    cp.title = 'Offer letter, Delacroix'
    cp.comments = 'v3, after legal review'
    cp.revision = 14
    cp.category = 'HR'
    cp.keywords = 'offer; confidential'
    d.save(os.path.join(TST, 'offer.docx'))
    add_custom_and_app(os.path.join(TST, 'offer.docx'), 'Microsoft Office Word', 'Delacroix Partners', 'H. Delacroix', 847)

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'Client'
    ws['B1'] = 'Fee'
    ws['A2'] = 'Northwind'
    ws['B2'] = 4200
    wb.properties.creator = 'Priya Raman'
    wb.properties.lastModifiedBy = 'H. Delacroix'
    wb.properties.title = 'Q2 fees'
    wb.properties.revision = 6
    wb.save(os.path.join(TST, 'fees.xlsx'))
    add_custom_and_app(os.path.join(TST, 'fees.xlsx'), 'Microsoft Excel', 'Delacroix Partners', '', 212)

    from pptx import Presentation
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = 'Pitch, internal only'
    p.core_properties.author = 'Priya Raman'
    p.core_properties.last_modified_by = 'H. Delacroix'
    p.core_properties.revision = 9
    p.save(os.path.join(TST, 'pitch.pptx'))
    add_custom_and_app(os.path.join(TST, 'pitch.pptx'), 'Microsoft Office PowerPoint', 'Delacroix Partners', 'H. Delacroix', 63)

    # PDF with an Info dictionary and an uncompressed XMP stream.
    open(os.path.join(TST, 'memo.pdf'), 'wb').write(make_pdf())

    # Unsupported type.
    open(os.path.join(TST, 'notes.txt'), 'w').write('plain text has no metadata FileSanity reads\n')
    print('tests/files:', sorted(os.listdir(TST)))


def add_custom_and_app(path, app, company, manager, total_time):
    """Replace docProps/app.xml with one that carries Company, Manager and
    TotalTime, and add docProps/custom.xml. Real Office writes these."""
    z = zipfile.ZipFile(path)
    items = {n: z.read(n) for n in z.namelist()}
    z.close()
    items['docProps/app.xml'] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        f'<Application>{app}</Application><AppVersion>16.0000</AppVersion>'
        f'<Company>{company}</Company>' + (f'<Manager>{manager}</Manager>' if manager else '') +
        f'<TotalTime>{total_time}</TotalTime><Template>Normal.dotm</Template></Properties>'
    ).encode()
    items['docProps/custom.xml'] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="ClientMatter"><vt:lpwstr>DEL-2026-0417</vt:lpwstr></property>'
        '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="3" name="Reviewer"><vt:lpwstr>S. Byrne</vt:lpwstr></property>'
        '</Properties>'
    ).encode()
    ct = items['[Content_Types].xml'].decode()
    if 'custom.xml' not in ct:
        ct = ct.replace('</Types>', '<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>')
        items['[Content_Types].xml'] = ct.encode()
    rels = items['_rels/.rels'].decode()
    if 'custom.xml' not in rels:
        rels = rels.replace('</Relationships>', '<Relationship Id="rIdCustom" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>')
        items['_rels/.rels'] = rels.encode()
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as out:
        for n, b in items.items():
            out.writestr(n, b)


def make_pdf():
    xmp = XMP.replace(b'Adobe Lightroom 8.2 (iOS)', b'Microsoft Word for Mac  ')
    objs = []
    objs.append(b'<< /Type /Catalog /Pages 2 0 R /Metadata 6 0 R >>')
    objs.append(b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>')
    objs.append(b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>')
    content = b'BT /F1 18 Tf 72 760 Td (Internal memo. Do not forward.) Tj ET'
    objs.append(b'<< /Length ' + str(len(content)).encode() + b' >>\nstream\n' + content + b'\nendstream')
    objs.append(b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>')
    objs.append(b'<< /Type /Metadata /Subtype /XML /Length ' + str(len(xmp)).encode() + b' >>\nstream\n' + xmp + b'\nendstream')
    objs.append(b'<< /Title (Internal memo) /Author (Priya Raman) /Subject (Northwind fees) /Keywords (confidential) '
                b'/Creator (Microsoft Word for Mac) /Producer (macOS Version 15.4 Quartz PDFContext) '
                b'/CreationDate (D:20260412090351+01\'00\') /ModDate (D:20260412091200+01\'00\') >>')
    out = bytearray(b'%PDF-1.4\n%\xe2\xe3\xcf\xd3\n')
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f'{i} 0 obj\n'.encode() + o + b'\nendobj\n'
    xref = len(out)
    out += f'xref\n0 {len(objs) + 1}\n'.encode() + b'0000000000 65535 f \n'
    for off in offsets:
        out += f'{off:010d} 00000 n \n'.encode()
    out += f'trailer\n<< /Size {len(objs) + 1} /Root 1 0 R /Info 7 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode()
    return bytes(out)


if __name__ == '__main__':
    main()
